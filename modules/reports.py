"""
modules/reports.py — Excel export using pandas + openpyxl.

Generates formatted .xlsx reports for sessions, revenue, and vehicles.
Header rows are bold + coloured; columns auto-sized.

Why pandas + openpyxl: pandas simplifies DataFrame construction and to_excel();
openpyxl's ExcelWriter gives us full control over header formatting and
column widths — a specific SRS requirement (FR-19).
"""

import os
from datetime import datetime
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from database.db import query_db

EXPORTS_DIR = os.path.join(os.path.dirname(os.path.dirname(__file__)), "exports")


def _auto_size_columns(ws) -> None:
    """Adjust each column's width to fit its longest cell value."""
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            try:
                cell_len = len(str(cell.value)) if cell.value else 0
                max_len = max(max_len, cell_len)
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max_len + 4, 60)


def _style_header(ws, header_color: str = "1B2A3B") -> None:
    """Bold + coloured background for the first (header) row."""
    for cell in ws[1]:
        cell.font      = Font(bold=True, color="FFFFFF")
        cell.fill      = PatternFill("solid", fgColor=header_color)
        cell.alignment = Alignment(horizontal="center")


def export_sessions(filename: str = None) -> str:
    """Export all parking sessions to a formatted .xlsx file.

    Returns:
        Absolute path of the created file.
    """
    rows = query_db(
        """SELECT ps.session_id, v.plate_no, v.owner_name, v.vehicle_type,
                  s.slot_code, s.zone,
                  ps.entry_time, ps.exit_time,
                  ps.fee, ps.status
           FROM parking_sessions ps
           JOIN vehicles v ON ps.vehicle_id = v.vehicle_id
           JOIN slots    s ON ps.slot_id    = s.slot_id
           ORDER BY ps.entry_time DESC"""
    )
    df = pd.DataFrame(rows)
    if df.empty:
        df = pd.DataFrame(columns=[
            "session_id","plate_no","owner_name","vehicle_type",
            "slot_code","zone","entry_time","exit_time","fee","status"
        ])
    df.columns = [c.replace("_", " ").title() for c in df.columns]

    if not filename:
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = os.path.join(EXPORTS_DIR, f"sessions_{ts}.xlsx")

    os.makedirs(EXPORTS_DIR, exist_ok=True)
    with pd.ExcelWriter(filename, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Sessions", index=False)
        ws = writer.sheets["Sessions"]
        _style_header(ws)
        _auto_size_columns(ws)

    return filename


def export_revenue(filename: str = None) -> str:
    """Export daily revenue summary to .xlsx."""
    rows = query_db(
        """SELECT date(exit_time) AS Date,
                  COUNT(*) AS Sessions,
                  ROUND(SUM(fee), 2) AS "Revenue (INR)"
           FROM parking_sessions
           WHERE status='Completed' AND exit_time IS NOT NULL
           GROUP BY Date
           ORDER BY Date DESC"""
    )
    df = pd.DataFrame(rows) if rows else pd.DataFrame(columns=["Date","Sessions","Revenue (INR)"])

    if not filename:
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = os.path.join(EXPORTS_DIR, f"revenue_{ts}.xlsx")

    os.makedirs(EXPORTS_DIR, exist_ok=True)
    with pd.ExcelWriter(filename, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Revenue", index=False)
        ws = writer.sheets["Revenue"]
        _style_header(ws, header_color="0F2613")
        _auto_size_columns(ws)

    return filename


def export_vehicles(filename: str = None) -> str:
    """Export the full vehicle register to .xlsx."""
    rows = query_db("SELECT * FROM vehicles ORDER BY created_at DESC")
    df = pd.DataFrame(rows) if rows else pd.DataFrame(
        columns=["vehicle_id","plate_no","owner_name","phone","vehicle_type","created_at"]
    )
    df.columns = [c.replace("_", " ").title() for c in df.columns]

    if not filename:
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = os.path.join(EXPORTS_DIR, f"vehicles_{ts}.xlsx")

    os.makedirs(EXPORTS_DIR, exist_ok=True)
    with pd.ExcelWriter(filename, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Vehicles", index=False)
        ws = writer.sheets["Vehicles"]
        _style_header(ws, header_color="2B1A0F")
        _auto_size_columns(ws)

    return filename
