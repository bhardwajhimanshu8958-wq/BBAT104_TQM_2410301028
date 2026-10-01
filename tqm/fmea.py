"""
tqm/fmea.py — Failure Mode and Effects Analysis (FMEA) module.

Manages the FMEA risk matrix: computing RPN (Severity × Occurrence × Detection),
providing risk-ranked views, and supporting Excel export.

Why FMEA? FMEA is a proactive TQM risk analysis tool that identifies failure
modes before they occur in production, reducing defect escape rates.
RPN = S × O × D; threshold for immediate mitigation: RPN > 100.
"""

import streamlit as st
import pandas as pd
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from database.db import query_db, execute_db
from modules.reports import EXPORTS_DIR
from openpyxl.styles import Font, PatternFill, Alignment
import os


def get_all_fmea() -> list[dict]:
    """Return all FMEA entries sorted by RPN descending."""
    return query_db("SELECT * FROM fmea ORDER BY rpn DESC")


def compute_rpn(severity: int, occurrence: int, detection: int) -> int:
    """Compute Risk Priority Number. RPN = Severity × Occurrence × Detection."""
    return severity * occurrence * detection


def upsert_fmea_row(process_step: str, failure_mode: str, effect: str, cause: str,
                     severity: int, occurrence: int, detection: int, mitigation: str,
                     status: str = "Open", fmea_id: int = None) -> int:
    """Insert a new FMEA row or update an existing one. Returns the fmea_id."""
    rpn = compute_rpn(severity, occurrence, detection)
    if fmea_id:
        execute_db(
            """UPDATE fmea SET process_step=?, failure_mode=?, effect=?, cause=?,
               severity=?, occurrence=?, detection=?, rpn=?, mitigation=?, status=?
               WHERE fmea_id=?""",
            (process_step, failure_mode, effect, cause, severity, occurrence,
             detection, rpn, mitigation, status, fmea_id),
        )
        return fmea_id
    else:
        return execute_db(
            """INSERT INTO fmea (process_step, failure_mode, effect, cause,
               severity, occurrence, detection, rpn, mitigation, status)
               VALUES (?,?,?,?,?,?,?,?,?,?)""",
            (process_step, failure_mode, effect, cause, severity, occurrence,
             detection, rpn, mitigation, status),
        )


def seed_fmea_if_empty() -> None:
    """Populate the FMEA table with 12 pre-defined failure modes if empty."""
    if query_db("SELECT 1 FROM fmea LIMIT 1"):
        return  # Already seeded

    fmea_rows = [
        # (process_step, failure_mode, effect, cause, S, O, D, mitigation)
        ("Vehicle Entry",      "Invalid plate format accepted",       "Corrupt vehicle record; mis-billing",      "Regex validator not applied",           8, 4, 3, "Added validate_plate() Poka-Yoke before DB insert"),
        ("Vehicle Entry",      "Duplicate active session created",    "Two sessions for same vehicle; data corruption", "No session check before record_entry()", 9, 3, 2, "Added session pre-check in record_entry(); raises ValueError"),
        ("Slot Allocation",    "Wrong slot type assigned (EV→Car)",  "EV vehicle has no charger; CTQ failure",   "Type filter bug in get_available_slot()", 7, 3, 4, "Unit tested type filter; Bike/EV/Car separate tuple filters"),
        ("Vehicle Exit",       "Fee calculated as zero or negative",  "Revenue loss; wrong audit trail",          "Duration <= 0 due to same-second exit",  9, 2, 2, "raise ValueError if delta <= 0 in calculate_duration_minutes"),
        ("Vehicle Exit",       "Exit time before entry accepted",     "Negative duration; invalid receipt",       "No timestamp comparison guard",          8, 2, 3, "Added validate_exit_after_entry() Poka-Yoke check"),
        ("Billing",            "Free minutes not deducted correctly", "Overcharge; customer complaint",           "Off-by-one in free_minutes comparison",  6, 4, 5, "Explicit test: calculate_fee('Car',10)==20.0 (base rate only)"),
        ("Database",           "FK constraint not enforced on delete","Orphaned sessions; slot stuck Occupied",   "SQLite FK off by default",               8, 3, 3, "PRAGMA foreign_keys=ON in get_connection(); confirmed with test"),
        ("UI / Streamlit",     "Page crashes on empty DB state",      "App unusable on fresh install",            "Missing None check on query results",    7, 5, 4, "All query results guarded with 'if rows else default'"),
        ("Reports",            "Excel column width too narrow",       "Data truncated; unreadable export",        "_auto_size_columns() not called first",  3, 5, 6, "Auto-sizer called for every sheet in export functions"),
        ("Network / Push",     "git push fails mid-session",         "Commits lost; GitHub desync",              "Unstable Wi-Fi / firewall block on :443", 9, 4, 2, "Documented in ERROR_LOG.md; use mobile hotspot as fallback"),
        ("Slot Management",    "Occupied slot deleted",              "Session references deleted slot; FK error", "No occupancy check before DELETE",       8, 3, 2, "delete_slot() queries status; raises ValueError if Occupied"),
        ("Vehicle Registry",   "Phone stored as text not digits",    "Search by phone fails; Poka-Yoke gap",      "No phone normalization in create_vehicle", 4, 6, 5, "validate_phone() normalizes and validates before insert"),
    ]

    for row in fmea_rows:
        ps, fm, eff, cause, s, o, d, mit = row
        rpn = s * o * d
        execute_db(
            """INSERT INTO fmea (process_step, failure_mode, effect, cause,
               severity, occurrence, detection, rpn, mitigation, status)
               VALUES (?,?,?,?,?,?,?,?,?,?)""",
            (ps, fm, eff, cause, s, o, d, rpn, mit, "Resolved"),
        )


def export_fmea_excel() -> str:
    """Export FMEA matrix to a formatted .xlsx file."""
    rows = get_all_fmea()
    df = pd.DataFrame(rows) if rows else pd.DataFrame()

    os.makedirs(EXPORTS_DIR, exist_ok=True)
    from datetime import datetime
    fname = os.path.join(EXPORTS_DIR, f"fmea_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx")

    with pd.ExcelWriter(fname, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="FMEA", index=False)
        ws = writer.sheets["FMEA"]
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="7B2CBF")
            cell.alignment = Alignment(horizontal="center")
        for col in ws.columns:
            max_len = max((len(str(c.value)) for c in col if c.value), default=10)
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 50)

    return fname


def render_fmea() -> None:
    """Render the FMEA risk matrix page with sortable table, RPN chart, and export."""
    seed_fmea_if_empty()

    st.title("⚠️ Failure Mode and Effects Analysis (FMEA)")
    st.caption("Proactive risk identification using RPN = Severity × Occurrence × Detection.")
    st.markdown(
        """
        > **Rating Rubric**: S (1=minor inconvenience, 10=system shutdown), O (1=remote chance, 10=almost certain), D (1=will catch immediately, 10=undetectable)  
        > **Action threshold**: RPN > 100 requires immediate corrective action.
        """
    )

    rows = get_all_fmea()
    if not rows:
        st.info("FMEA table is being populated.")
        return

    tab_matrix, tab_chart, tab_export = st.tabs(["📊 Risk Matrix", "📉 RPN Risk Chart", "📥 Excel Export"])

    with tab_matrix:
        df = pd.DataFrame(rows)
        df["Risk Level"] = df["rpn"].apply(lambda x: "🔴 Critical" if x > 100 else ("🟡 Moderate" if x > 50 else "🟢 Low"))
        st.dataframe(
            df,
            column_config={
                "fmea_id": "ID",
                "process_step": "Process Step",
                "failure_mode": "Failure Mode",
                "effect": "Potential Effect",
                "cause": "Root Cause",
                "severity": st.column_config.NumberColumn("S (1-10)"),
                "occurrence": st.column_config.NumberColumn("O (1-10)"),
                "detection": st.column_config.NumberColumn("D (1-10)"),
                "rpn": st.column_config.NumberColumn("RPN (S×O×D)"),
                "Risk Level": "Risk Level",
                "mitigation": "Mitigation Action",
                "status": "Status",
            },
            hide_index=True,
            use_container_width=True,
        )

        top5 = df.head(5)
        st.subheader("🚨 Top 5 High-Priority Failure Modes")
        for _, r in top5.iterrows():
            risk = "🔴 Critical" if r["rpn"] > 100 else "🟡 Moderate"
            with st.expander(f"{risk} RPN {int(r['rpn'])} — {r['failure_mode']}", expanded=False):
                c1, c2 = st.columns(2)
                c1.write(f"**Process**: {r['process_step']}")
                c1.write(f"**Effect**: {r['effect']}")
                c1.write(f"**Cause**: {r['cause']}")
                c2.write(f"**S × O × D**: {int(r['severity'])} × {int(r['occurrence'])} × {int(r['detection'])} = **{int(r['rpn'])}**")
                c2.write(f"**Mitigation**: {r['mitigation']}")
                c2.write(f"**Status**: {r['status']}")

    with tab_chart:
        df = pd.DataFrame(rows)
        fig, ax = plt.subplots(figsize=(10, 5), facecolor="#0d1b2a")
        ax.set_facecolor("#1b2a3b")
        colors = ["#e63946" if x > 100 else ("#f4a261" if x > 50 else "#2a9d8f") for x in df["rpn"]]
        bars = ax.barh(df["failure_mode"], df["rpn"], color=colors, alpha=0.9)
        ax.axvline(100, color="#ffffff", linestyle="--", linewidth=1.5, label="RPN=100 Action Threshold")
        ax.set_xlabel("Risk Priority Number (RPN)", color="#aaa")
        ax.tick_params(colors="#ccc", labelsize=8)
        ax.spines[:].set_color("#333")
        ax.legend(fontsize=9, labelcolor="#fff", facecolor="none", edgecolor="#555")
        plt.title("FMEA Risk Priority Numbers — Parking Management System", color="#e0f4ff", fontsize=11, fontweight="bold", pad=12)
        plt.tight_layout()
        st.pyplot(fig, use_container_width=True)
        plt.close(fig)

    with tab_export:
        st.subheader("Export FMEA Matrix to Excel")
        if st.button("Generate FMEA Excel Workbook", use_container_width=True):
            with st.spinner("Generating..."):
                fpath = export_fmea_excel()
                with open(fpath, "rb") as f:
                    st.download_button("⬇️ Download FMEA.xlsx", data=f.read(),
                                      file_name=os.path.basename(fpath),
                                      mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                      use_container_width=True)
                st.success("Workbook ready for download!")
